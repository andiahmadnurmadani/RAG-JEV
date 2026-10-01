"""Buat contoh berkas asli (Word/Excel/PowerPoint/ODT/ePub/RTF/teks) untuk menguji pembacaan.

Kenapa ada skrip terpisah: service ini sengaja hanya butuh pustaka standar untuk membaca
berkas, jadi tidak ada python-docx/openpyxl di venv service. Skrip ini dijalankan dengan venv
sementara yang punya pustaka aslinya, supaya berkas ujinya benar-benar dibuat oleh pustaka
Office (bukan struktur zip karangan sendiri).

    uv venv data/tmp/genvenv --python 3.11
    uv pip install --python data/tmp/genvenv/Scripts/python.exe python-docx openpyxl python-pptx odfpy ebooklib
    data/tmp/genvenv/Scripts/python.exe scripts/make_format_samples.py data/tmp/fmt
"""

from __future__ import annotations

import sys
from pathlib import Path

MARKERS = {
    "docx": "KUA2026DOCX",
    "xlsx": "KUA2026XLSX",
    "pptx": "KUA2026NOTES",
    "odt": "KUA2026ODT",
    "epub": "KUA2026EPUB",
}


def write_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Kebijakan Cuti 2026", level=1)
    document.add_paragraph(
        "Kuota cuti tahunan karyawan tetap adalah 12 hari kerja. Kafe kantor buka pukul 08.00 "
        f"hingga 17.00. Kode kebijakan {MARKERS['docx']} berlaku sejak Januari."
    )
    document.add_paragraph("Ringkasan aturan:")
    table = document.add_table(rows=3, cols=3)
    rows = [
        ("Bagian", "Kuota", "Catatan"),
        ("Keuangan", "12", "dibulatkan per tahun"),
        ("Operasi", "12", "ganti tahun tidak menumpuk"),
    ]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    document.add_paragraph("Pengesahan: Departemen SDM, café internal menjadi lampiran.")
    document.save(str(path))


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kuota"
    sheet.append(["Nama", "Bagian", "Kuota Cuti"])
    sheet.append(["Andi", "Keuangan", 12])
    sheet.append(["Budi", "Operasi", 12])
    sheet.append(["Kode kebijakan", MARKERS["xlsx"], None])
    sheet["E1"] = "Total"
    sheet["E2"] = "=SUM(C2:C3)"
    notes = workbook.create_sheet("Catatan")
    notes.append(["Catatan rapat", "Lembur dibayar setelah jam kerja"])
    notes.append(["Lampiran", "café internal 08.00-17.00"])
    workbook.save(str(path))


def write_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Pemeliharaan Rutin"
    slide.placeholders[1].text = "Kuota cuti dua belas hari\nLaporan bulanan setiap tanggal lima"
    notes = slide.notes_slide.notes_text_frame
    notes.text = f"Catatan pembicara: sebut kode {MARKERS['pptx']} dan lampiran café."

    second = presentation.slides.add_slide(presentation.slide_layouts[1])
    second.shapes.title.text = "Langkah Berikutnya"
    second.placeholders[1].text = "Evaluasi kuota kuartal depan"
    presentation.save(str(path))


def write_odt(path: Path) -> None:
    from odf.opendocument import OpenDocumentText
    from odf.table import Table, TableCell, TableRow
    from odf.text import H, P

    document = OpenDocumentText()
    document.text.addElement(H(outlinelevel=1, text="Kebijakan Cuti"))
    document.text.addElement(
        P(text=f"Kuota cuti tahunan karyawan tetap 12 hari. Kode kebijakan {MARKERS['odt']}.")
    )
    table = Table(name="Kuota")
    for values in (("Bagian", "Kuota"), ("Keuangan", "12"), ("Operasi", "12")):
        row = TableRow()
        for value in values:
            cell = TableCell()
            cell.addElement(P(text=value))
            row.addElement(cell)
        table.addElement(row)
    document.text.addElement(table)
    document.save(str(path))


def write_epub(path: Path) -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("rag-service-uji-format")
    book.set_title("Buku Uji Format")
    book.set_language("id")
    chapters = []
    for number, text in enumerate(
        (
            f"Bab Satu. Cuti tahunan dua belas hari kerja. Kode {MARKERS['epub']}.",
            "Bab Dua. Lembur dibayar setelah jam kerja dan harus disetujui atasan.",
        ),
        start=1,
    ):
        chapter = epub.EpubHtml(title=f"Bab {number}", file_name=f"bab{number}.xhtml", lang="id")
        chapter.content = f"<h1>Bab {number}</h1><p>{text}</p>"
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", *chapters]
    epub.write_epub(str(path), book)


def write_plain(path: Path, text: str, encoding: str = "utf-8") -> None:
    path.write_bytes(text.encode(encoding))


def write_rtf(path: Path) -> None:
    body = (
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}{\colortbl ;\red0\green0\blue0;}"
        r"\f0\fs24 SOP Cuti 2026\par"
        r"Kuota cuti adalah 12 hari kerja. Kode kebijakan KUA2026RTF berlaku dari Januari.\par"
        r"Biaya kafe \'e9 00 ditanggung kantor.\par}"
    )
    path.write_bytes(body.encode("latin-1"))


def main() -> int:
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "data/tmp/fmt")
    target.mkdir(parents=True, exist_ok=True)
    write_docx(target / "kebijakan_cuti.docx")
    write_xlsx(target / "kuota_cuti.xlsx")
    write_pptx(target / "rapat_rutin.pptx")
    write_odt(target / "kebijakan_cuti.odt")
    write_epub(target / "buku_uji.epub")
    write_rtf(target / "sop_cuti.rtf")
    write_plain(
        target / "sop_cuti.md",
        "# SOP Cuti\n\nKuota cuti 12 hari. Kode KUA2026MD. Tabel:\n\n| Bagian | Kuota |\n|---|---|\n| Keuangan | 12 |\n",
    )
    write_plain(target / "catatan_latin.txt", "Ringkasan: kuota cuti 12 hari, café buka 08.00. Kode KUA2026LATIN.\n", "latin-1")
    write_plain(target / "kuota.csv", "nama;bagian;kuota\nAndi;Keuangan;12\nBudi;Operasi;12\nKode;KUA2026CSV;-\n")
    write_plain(target / "kebijakan.yaml", "kebijakan: cuti\nkode: KUA2026YAML\nkuota: 12\n")
    write_plain(target / "sop.html", "<html><head><title>SOP Cuti</title></head><body><h1>SOP Cuti</h1>"
                                     "<p>Kode KUA2026HTML. Kuota 12 hari.</p><table><tr><td>Bagian</td><td>Kuota</td></tr>"
                                     "<tr><td>Keuangan</td><td>12</td></tr></table></body></html>")
    write_plain(target / "audit.log", "2026-09-30 INFO mulai\n2026-09-30 ERROR gagal simpan KUA2026LOG\n")
    print("berkas contoh di:", target.resolve())
    for item in sorted(target.iterdir()):
        print(f"  {item.name:24} {item.stat().st_size:>7} byte")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
